import os
import tempfile
from unittest.mock import patch, MagicMock
from django.test import TestCase
from django.core.files.uploadedfile import SimpleUploadedFile
import docx

from rest_framework.test import APIClient
from rest_framework import status

from apps.authentication.models import User
from apps.cv.models import CandidateCV, ParsedCV
from apps.cv.services.contracts import EducationItem, ExperienceItem, ExtractionResult
from apps.cv.services.text_extractor import (
    clean_extracted_text,
    extract_text_from_docx,
    extract_text_from_pdf,
    extract_text,
)
from apps.cv.services.skill_extractor import SkillExtractor, extract_skills, get_skill_extractor
from apps.cv.services.entity_extractor import (
    EntityExtractor,
    extract_education,
    extract_experience,
    extract_entities,
    get_entity_extractor,
)
from apps.cv.services.pipeline import CVExtractionPipeline, extract_cv_data


class ContractsTests(TestCase):
    def test_education_item(self):
        item = EducationItem(degree='BSc Computer Science', institution='MIT', year='2024', grade='3.9')
        d = item.to_dict()
        self.assertEqual(d['degree'], 'BSc Computer Science')
        self.assertEqual(d['institution'], 'MIT')
        self.assertEqual(d['year'], '2024')
        self.assertEqual(d['grade'], '3.9')

    def test_experience_item(self):
        item = ExperienceItem(title='Software Engineer', company='Acme', duration='2 years', years_count=2.0)
        d = item.to_dict()
        self.assertEqual(d['title'], 'Software Engineer')
        self.assertEqual(d['company'], 'Acme')
        self.assertEqual(d['years_count'], 2.0)

    def test_extraction_result(self):
        res = ExtractionResult(raw_text='Test resume', skills=['Python'], word_count=2)
        d = res.to_dict()
        self.assertEqual(d['raw_text'], 'Test resume')
        self.assertEqual(d['skills'], ['Python'])
        self.assertEqual(d['word_count'], 2)
        self.assertTrue(d['success'])
        self.assertIsNone(d['error_message'])


class TextExtractorTests(TestCase):
    def test_clean_extracted_text_normalizes_whitespace_and_bullets(self):
        dirty = "\u2022 Item 1  with   spaces\n\u2023 Item 2\n\n\n\n\nItem 3 with non-breaking\u00a0space"
        cleaned = clean_extracted_text(dirty)
        self.assertIn('- Item 1 with spaces', cleaned)
        self.assertIn('- Item 2', cleaned)
        self.assertIn('Item 3 with non-breaking space', cleaned)
        self.assertNotIn('\n\n\n', cleaned)

    def test_clean_extracted_text_empty(self):
        self.assertEqual(clean_extracted_text(None), '')
        self.assertEqual(clean_extracted_text(''), '')

    def test_docx_text_extraction_paragraphs_and_tables(self):
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
            docx_path = f.name

        try:
            doc = docx.Document()
            doc.add_heading('John Candidate', 0)
            doc.add_paragraph('Full-stack developer with 5 years experience.')
            table = doc.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = 'Skills: Python, Django'
            table.rows[0].cells[1].text = 'Location: Dhaka'
            doc.save(docx_path)

            extracted = extract_text_from_docx(docx_path)
            self.assertIn('John Candidate', extracted)
            self.assertIn('Full-stack developer with 5 years experience.', extracted)
            self.assertIn('Skills: Python, Django', extracted)
            self.assertIn('Location: Dhaka', extracted)
        finally:
            if os.path.exists(docx_path):
                os.unlink(docx_path)

    @patch('pdfplumber.open')
    def test_pdf_text_extraction_with_mock(self, mock_open):
        mock_pdf = MagicMock()
        mock_page1 = MagicMock()
        mock_page1.extract_text.return_value = 'Page 1: Jane Doe, Software Engineer'
        mock_page2 = MagicMock()
        mock_page2.extract_text.return_value = 'Page 2: Education: Stanford University'
        mock_pdf.pages = [mock_page1, mock_page2]
        mock_pdf.__enter__.return_value = mock_pdf

        mock_open.return_value = mock_pdf

        with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as f:
            pdf_path = f.name

        try:
            extracted = extract_text_from_pdf(pdf_path)
            self.assertIn('Page 1: Jane Doe, Software Engineer', extracted)
            self.assertIn('Page 2: Education: Stanford University', extracted)
        finally:
            if os.path.exists(pdf_path):
                os.unlink(pdf_path)

    def test_extract_text_file_not_found(self):
        with self.assertRaises(FileNotFoundError):
            extract_text('non_existent_file.pdf')

    def test_extract_text_unsupported_format(self):
        with tempfile.NamedTemporaryFile(suffix='.txt', delete=False) as f:
            f.write(b'sample text')
            txt_path = f.name

        try:
            with self.assertRaises(ValueError):
                extract_text(txt_path)
        finally:
            if os.path.exists(txt_path):
                os.unlink(txt_path)


class SkillExtractorTests(TestCase):
    def setUp(self):
        self.extractor = SkillExtractor()

    def test_extract_standard_tech_skills(self):
        text = 'Extensive experience building web applications using Python, Django, React, and PostgreSQL.'
        skills = self.extractor.extract_skills(text)
        self.assertIn('Python', skills)
        self.assertIn('Django', skills)
        self.assertIn('React', skills)
        self.assertIn('PostgreSQL', skills)

    def test_extract_symbol_skills(self):
        text = 'Strong knowledge in C++, C#, .NET Core, Node.js, and CI/CD automation.'
        skills = self.extractor.extract_skills(text)
        self.assertIn('C++', skills)
        self.assertIn('C#', skills)
        self.assertIn('.NET Core', skills)
        self.assertIn('Node.js', skills)
        self.assertIn('CI/CD', skills)

    def test_canonical_synonyms_and_aliases(self):
        text = 'Proficient in React.js, Golang, Postgres, and Tailwind CSS.'
        skills = self.extractor.extract_skills(text)
        self.assertIn('React', skills)
        self.assertIn('Go', skills)
        self.assertIn('PostgreSQL', skills)
        self.assertIn('TailwindCSS', skills)

    def test_avoid_false_positives_for_single_letter_skills(self):
        text = 'Candidate was clear, created clever code for the company car club.'
        skills = self.extractor.extract_skills(text)
        # Should NOT match single letter 'C' from words starting with 'c'
        self.assertNotIn('C', skills)

        # But should match when explicitly listed as language
        text_with_c = 'Programming languages: Python, C, C++, Java'
        skills_with_c = self.extractor.extract_skills(text_with_c)
        self.assertIn('C', skills_with_c)
        self.assertIn('Python', skills_with_c)
        self.assertIn('C++', skills_with_c)
        self.assertIn('Java', skills_with_c)

    def test_extract_skills_by_category(self):
        text = 'Skilled in Python, Docker, Redis, and Machine Learning.'
        by_category = self.extractor.extract_skills_by_category(text)
        self.assertIn('programming_languages', by_category)
        self.assertIn('Python', by_category['programming_languages'])
        self.assertIn('cloud_and_devops', by_category)
        self.assertIn('Docker', by_category['cloud_and_devops'])
        self.assertIn('databases', by_category)
        self.assertIn('Redis', by_category['databases'])
        self.assertIn('data_science_and_ai', by_category)
        self.assertIn('Machine Learning', by_category['data_science_and_ai'])

    def test_extract_skills_with_counts(self):
        text = 'Python developer with strong Python backend and Python script automation using Django.'
        counts = self.extractor.extract_skills_with_counts(text)
        self.assertEqual(counts.get('Python'), 3)
        self.assertEqual(counts.get('Django'), 1)


class CVExtractionPipelineTests(TestCase):
    def setUp(self):
        self.student = User.objects.create_user(
            email='bonny_dev@careerflow.demo',
            username='bonny_dev',
            password='Password123!',
            role=User.Role.STUDENT,
            full_name='Bonny Developer'
        )

    def test_pipeline_with_docx_file(self):
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
            docx_path = f.name

        try:
            doc = docx.Document()
            doc.add_heading('Bonny - Backend Engineer', 0)
            doc.add_paragraph('Expert in Python, FastAPI, Docker, and PostgreSQL.')
            doc.save(docx_path)

            pipeline = CVExtractionPipeline()
            result = pipeline.process_file(docx_path)

            self.assertTrue(result.success)
            self.assertGreater(result.word_count, 0)
            self.assertIn('Python', result.skills)
            self.assertIn('FastAPI', result.skills)
            self.assertIn('Docker', result.skills)
            self.assertIn('PostgreSQL', result.skills)
        finally:
            if os.path.exists(docx_path):
                os.unlink(docx_path)

    def test_pipeline_with_candidate_cv_model(self):
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
            docx_path = f.name

        try:
            doc = docx.Document()
            doc.add_heading('Model Candidate', 0)
            doc.add_paragraph('Experienced with Django, Redis, and Git.')
            doc.save(docx_path)

            with open(docx_path, 'rb') as f_read:
                cv_file = SimpleUploadedFile('test_resume.docx', f_read.read(), content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

            candidate_cv = CandidateCV.objects.create(
                user=self.student,
                file=cv_file,
                original_filename='test_resume.docx',
                file_size=1024,
                file_type='docx',
                is_active=True
            )

            pipeline = CVExtractionPipeline()
            result = pipeline.process_candidate_cv(candidate_cv)

            self.assertTrue(result.success)
            self.assertIn('Django', result.skills)
            self.assertIn('Redis', result.skills)
            self.assertIn('Git', result.skills)
        finally:
            if os.path.exists(docx_path):
                os.unlink(docx_path)

    def test_pipeline_with_pluggable_entity_extractor(self):
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
            docx_path = f.name

        try:
            doc = docx.Document()
            doc.add_paragraph('Sample CV with education and experience.')
            doc.save(docx_path)

            def mock_person2_entities(text):
                return {
                    'education': [{'degree': 'BSc Software Engineering', 'institution': 'BRAC University'}],
                    'experience': [{'title': 'Junior Dev', 'company': 'InnoTech'}]
                }

            pipeline = CVExtractionPipeline(entity_extractor=mock_person2_entities)
            result = pipeline.process_file(docx_path)

            self.assertTrue(result.success)
            self.assertEqual(len(result.education), 1)
            self.assertEqual(result.education[0]['degree'], 'BSc Software Engineering')
            self.assertEqual(len(result.experience), 1)
            self.assertEqual(result.experience[0]['company'], 'InnoTech')
        finally:
            if os.path.exists(docx_path):
                os.unlink(docx_path)

    def test_pipeline_file_not_found(self):
        pipeline = CVExtractionPipeline()
        result = pipeline.process_file('imaginary_non_existent.pdf')
        self.assertFalse(result.success)
        self.assertIn('File not found', result.error_message)

    def test_extractor_facade_service_imports(self):
        from apps.cv.services.extractor import (
            CVExtractor,
            extract_text,
            extract_skills,
            extract_education,
            extract_experience,
            extract_cv_data,
        )
        extractor = CVExtractor()
        self.assertIsNotNone(extractor)
        self.assertTrue(callable(extract_text))
        self.assertTrue(callable(extract_skills))
        self.assertTrue(callable(extract_education))
        self.assertTrue(callable(extract_experience))
        self.assertTrue(callable(extract_cv_data))


class EntityExtractorTests(TestCase):
    """
    US-07-T4: Tests for Education and Experience extraction logic.
    """
    def setUp(self):
        self.extractor = EntityExtractor()

    def test_extract_education_degrees_and_institutions(self):
        text = """
        EDUCATION
        Bachelor of Science in Computer Science and Engineering
        BRAC University, 2019 - 2023
        CGPA: 3.85 / 4.0

        Master of Science in Software Engineering
        University of Oxford, 2024
        """
        edu = self.extractor.extract_education(text)
        self.assertGreaterEqual(len(edu), 2)
        degrees = [item['degree'].lower() for item in edu]
        self.assertTrue(any('bachelor' in d or 'science' in d for d in degrees))
        self.assertTrue(any('master' in d for d in degrees))

        institutions = [item['institution'] for item in edu]
        self.assertTrue(any('BRAC' in inst or 'University' in inst for inst in institutions))
        self.assertTrue(any('Oxford' in inst or 'University' in inst for inst in institutions))

    def test_extract_education_bba_mba_phd_diploma(self):
        text = """
        ACADEMICS:
        - MBA in Finance, Harvard Business School, 2021
        - PhD in Computer Science, Stanford University, 2018
        - Diploma in Information Technology, 2015
        """
        edu = self.extractor.extract_education(text)
        self.assertGreaterEqual(len(edu), 3)
        degrees = [item['degree'].lower() for item in edu]
        self.assertTrue(any('mba' in d for d in degrees))
        self.assertTrue(any('phd' in d or 'ph.d' in d for d in degrees))
        self.assertTrue(any('diploma' in d for d in degrees))

    def test_extract_experience_roles_and_companies(self):
        text = """
        WORK EXPERIENCE:
        Senior Software Engineer at Google
        Jan 2021 - Present
        - Architected cloud distributed microservices.

        Frontend Developer | InnoTech Solutions
        2018 - 2021 (3 years)
        - Developed React client interfaces.
        """
        exp = self.extractor.extract_experience(text)
        self.assertGreaterEqual(len(exp), 2)

        titles = [item['title'].lower() for item in exp]
        self.assertTrue(any('senior software engineer' in t for t in titles))
        self.assertTrue(any('frontend developer' in t for t in titles))

        companies = [item['company'] for item in exp]
        self.assertTrue(any('Google' in c for c in companies))
        self.assertTrue(any('InnoTech' in c for c in companies))

    def test_extract_experience_duration_years(self):
        text = """
        Backend Developer at Acme Corp
        2020 - 2024 (4 years)
        """
        exp = self.extractor.extract_experience(text)
        self.assertEqual(len(exp), 1)
        self.assertIn('duration_years', exp[0])
        self.assertEqual(exp[0]['duration_years'], 4.0)
        self.assertEqual(exp[0]['years_count'], 4.0)

    def test_empty_or_whitespace_text_extraction(self):
        self.assertEqual(self.extractor.extract_education(""), [])
        self.assertEqual(self.extractor.extract_experience("   \n\t  "), [])


class ParsedCVModelTests(TestCase):
    """
    US-07-T5: Tests for ParsedCV database schema and relationships.
    """
    def setUp(self):
        self.student = User.objects.create_user(
            email='parsed_student@careerflow.demo',
            username='parsed_student',
            password='Password123!',
            role=User.Role.STUDENT,
            full_name='Parsed Model Tester'
        )
        self.cv = CandidateCV.objects.create(
            user=self.student,
            file=SimpleUploadedFile('test.pdf', b'%PDF-1.4 test', content_type='application/pdf'),
            original_filename='test.pdf',
            file_size=100,
            file_type='pdf',
            is_active=True
        )

    def test_create_and_retrieve_parsed_cv(self):
        parsed = ParsedCV.objects.create(
            cv=self.cv,
            raw_text="Test CV with Python and Django",
            skills=["Python", "Django", "PostgreSQL"],
            education=[{"degree": "BSc CSE", "institution": "DU", "year": "2023"}],
            experience=[{"title": "Software Engineer", "company": "Tech", "duration_years": 2.0}]
        )
        self.assertEqual(parsed.cv, self.cv)
        self.assertEqual(parsed.skills_count, 3)
        self.assertIn("Python", parsed.skills)
        self.assertEqual(len(parsed.education), 1)
        self.assertEqual(len(parsed.experience), 1)
        self.assertEqual(parsed.cv.parsed_data, parsed)

    def test_cascade_delete_with_candidate_cv(self):
        ParsedCV.objects.create(
            cv=self.cv,
            raw_text="Candidate raw text",
            skills=["React"],
            education=[],
            experience=[]
        )
        self.assertEqual(ParsedCV.objects.count(), 1)
        self.cv.delete()
        self.assertEqual(ParsedCV.objects.count(), 0)


class CVExtractionEdgeCasesTests(TestCase):
    """
    US-07-T6: Tests for error handling, corrupted files, and OCR/scanned edge cases.
    """
    def setUp(self):
        self.pipeline = CVExtractionPipeline()

    def test_corrupted_pdf_file_handled_gracefully(self):
        with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as f:
            f.write(b'GARBAGE_CORRUPT_BYTES_NOT_A_REAL_PDF_1234567890')
            corrupt_pdf = f.name

        try:
            result = self.pipeline.process_file(corrupt_pdf)
            self.assertFalse(result.success)
            self.assertIsNotNone(result.error_message)
            self.assertTrue(
                'invalid' in result.error_message.lower() or
                'corrupt' in result.error_message.lower() or
                'pdf' in result.error_message.lower()
            )
        finally:
            if os.path.exists(corrupt_pdf):
                os.unlink(corrupt_pdf)

    def test_corrupted_docx_file_handled_gracefully(self):
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
            f.write(b'PK_FAKE_ZIP_HEADER_CORRUPTED_DOCX_FILE')
            corrupt_docx = f.name

        try:
            result = self.pipeline.process_file(corrupt_docx)
            self.assertFalse(result.success)
            self.assertIsNotNone(result.error_message)
            self.assertTrue(
                'invalid' in result.error_message.lower() or
                'corrupt' in result.error_message.lower() or
                'docx' in result.error_message.lower()
            )
        finally:
            if os.path.exists(corrupt_docx):
                os.unlink(corrupt_docx)

    @patch('apps.cv.services.pipeline.extract_text', return_value='')
    def test_scanned_or_empty_document_detection(self, mock_extract):
        with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as f:
            f.write(b'%PDF-1.4 dummy')
            dummy_pdf = f.name

        try:
            result = self.pipeline.process_file(dummy_pdf)
            self.assertTrue(result.success)
            self.assertTrue(result.is_scanned)
            self.assertIn('scanned', result.error_message.lower())
        finally:
            if os.path.exists(dummy_pdf):
                os.unlink(dummy_pdf)


class CVParseAPITests(TestCase):
    """
    US-07-T5 & US-07-T6: Tests for CV parsing API endpoints (/api/cv/<id>/parse/ and /api/cv/<id>/parsed/).
    """
    def setUp(self):
        self.client = APIClient()
        self.student = User.objects.create_user(
            email='api_student@careerflow.demo',
            username='api_student',
            password='Password123!',
            role=User.Role.STUDENT,
            full_name='API Student'
        )
        self.other_student = User.objects.create_user(
            email='other_student@careerflow.demo',
            username='other_student',
            password='Password123!',
            role=User.Role.STUDENT,
            full_name='Other Student'
        )
        self.client.force_authenticate(user=self.student)

        # Create valid sample docx for CV
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
            self.temp_docx = f.name
        doc = docx.Document()
        doc.add_heading('John Doe Resume', 0)
        doc.add_paragraph('Experienced Python, Django, Docker and PostgreSQL Developer.')
        doc.add_heading('Education', level=1)
        doc.add_paragraph('BSc in Computer Science, University of California, 2022')
        doc.add_heading('Experience', level=1)
        doc.add_paragraph('Software Engineer at TechCorp, 2022 - Present (2 years)')
        doc.save(self.temp_docx)

        with open(self.temp_docx, 'rb') as f_read:
            uploaded = SimpleUploadedFile(
                'resume.docx',
                f_read.read(),
                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            )

        self.cv = CandidateCV.objects.create(
            user=self.student,
            file=uploaded,
            original_filename='resume.docx',
            file_size=1024,
            file_type='docx',
            is_active=True
        )

    def tearDown(self):
        if os.path.exists(self.temp_docx):
            os.unlink(self.temp_docx)

    def test_parse_cv_endpoint_success(self):
        url = f"/api/cv/{self.cv.id}/parse/"
        response = self.client.post(url)
        self.assertIn(response.status_code, [status.HTTP_200_OK, status.HTTP_201_CREATED])
        self.assertIn('parsed_cv', response.data)
        data = response.data['parsed_cv']
        self.assertIn('Python', data['skills'])
        self.assertIn('Django', data['skills'])
        self.assertGreaterEqual(len(data['education']), 1)
        self.assertGreaterEqual(len(data['experience']), 1)

    def test_get_parsed_cv_endpoint_success(self):
        # Trigger parse first
        self.client.post(f"/api/cv/{self.cv.id}/parse/")
        url = f"/api/cv/{self.cv.id}/parsed/"
        response = self.client.get(url)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['cv'], self.cv.id)
        self.assertIn('skills', response.data)
        self.assertIn('education', response.data)
        self.assertIn('experience', response.data)

    def test_get_parsed_cv_not_found_when_unparsed(self):
        url = f"/api/cv/{self.cv.id}/parsed/"
        response = self.client.get(url)
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_parse_forbidden_for_other_student(self):
        self.client.force_authenticate(user=self.other_student)
        url = f"/api/cv/{self.cv.id}/parse/"
        response = self.client.post(url)
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_parse_non_existent_cv_returns_404(self):
        url = "/api/cv/999999/parse/"
        response = self.client.post(url)
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)


class ExtractionAccuracyBenchmarkTests(TestCase):
    """
    US-07-T7: Accuracy verification ensuring at least 80% skills identification on standard test resumes.
    Acceptance Criteria: 'Extracted skills list accurately identifies at least 80% of skills in standard test resumes.'
    """
    def setUp(self):
        self.extractor = SkillExtractor()

    def test_backend_resume_skills_accuracy_above_80_percent(self):
        resume_text = """
        Backend Software Engineer with 4 years experience building distributed systems.
        Proficient in Python, Django, REST APIs, PostgreSQL, Redis, Docker, and Celery.
        Strong expertise with Git version control, CI/CD pipelines, and Linux server environments.
        """
        ground_truth = ['Python', 'Django', 'PostgreSQL', 'Redis', 'Docker', 'Git', 'Linux']
        extracted = self.extractor.extract_skills(resume_text)
        extracted_set = set(extracted)

        matched = [s for s in ground_truth if s in extracted_set]
        recall = len(matched) / len(ground_truth)
        self.assertGreaterEqual(recall, 0.80, f"Backend resume skill recall {recall:.2f} was below 80% requirement!")

    def test_frontend_resume_skills_accuracy_above_80_percent(self):
        resume_text = """
        Frontend Web Developer specializing in single page applications.
        Core stack: JavaScript, TypeScript, React, Next.js, Redux, HTML5, CSS3, Tailwind CSS.
        Experience collaborating on UI design with Figma and using Jest for unit testing.
        """
        ground_truth = ['JavaScript', 'TypeScript', 'React', 'Next.js', 'Redux', 'HTML5', 'CSS3', 'TailwindCSS', 'Figma']
        extracted = self.extractor.extract_skills(resume_text)
        extracted_set = set(extracted)

        matched = [s for s in ground_truth if s in extracted_set]
        recall = len(matched) / len(ground_truth)
        self.assertGreaterEqual(recall, 0.80, f"Frontend resume skill recall {recall:.2f} was below 80% requirement!")

    def test_data_science_resume_skills_accuracy_above_80_percent(self):
        resume_text = """
        Data Scientist with background in machine learning and data analytics.
        Proficient in Python, SQL, Pandas, NumPy, Scikit-Learn, PyTorch, TensorFlow, and Deep Learning.
        Experienced with data visualization using Tableau and model deployment with Docker.
        """
        ground_truth = ['Python', 'SQL', 'Pandas', 'NumPy', 'Scikit-Learn', 'PyTorch', 'TensorFlow', 'Machine Learning', 'Docker']
        extracted = self.extractor.extract_skills(resume_text)
        extracted_set = set(extracted)

        matched = [s for s in ground_truth if s in extracted_set]
        recall = len(matched) / len(ground_truth)
        self.assertGreaterEqual(recall, 0.80, f"Data science skill recall {recall:.2f} was below 80% requirement!")

